"""Lecture des fichiers importés (CSV / TXT / XLSX) : détection du format, aperçu, lecture par blocs."""

import codecs
import csv
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from itertools import islice
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from app.services.errors import BusinessError

CSV_EXTENSIONS = {".csv", ".txt"}
EXCEL_EXTENSIONS = {".xlsx"}
SUPPORTED_EXTENSIONS = CSV_EXTENSIONS | EXCEL_EXTENSIONS
SNIFF_BYTES = 1 << 20
SEPARATORS = (";", ",", "\t")
PREVIEW_ROWS = 20


@dataclass(frozen=True)
class CsvFormat:
    encoding: str
    separator: str


def is_excel(path: Path) -> bool:
    return path.suffix.lower() in EXCEL_EXTENSIONS


def detect_csv_format(path: Path) -> CsvFormat:
    """Encodage (UTF-8 puis CP1252) et séparateur (; , tabulation) d'un fichier texte."""
    with path.open("rb") as handle:
        sample = handle.read(SNIFF_BYTES)
    try:
        codecs.getincrementaldecoder("utf-8")().decode(sample, final=False)
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        encoding = "cp1252"
    text = sample.decode(encoding, errors="ignore")
    header = text.splitlines()[0] if text else ""
    separator = max(SEPARATORS, key=header.count)
    if header.count(separator) == 0:
        try:
            separator = csv.Sniffer().sniff(text[:10000], delimiters="".join(SEPARATORS)).delimiter
        except csv.Error:
            separator = ";"
    return CsvFormat(encoding=encoding, separator=separator)


def excel_sheets(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def cell_to_str(value: Any) -> str:
    """Valeur de cellule Excel -> texte (entiers sans .0, dates ISO)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _unique_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, header in enumerate(headers):
        name = header.strip() or f"colonne_{index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        result.append(name)
    return result


def iter_chunks(path: Path, sheet_name: str | None, chunk_size: int) -> Iterator[pd.DataFrame]:
    """Blocs de lignes (toutes les valeurs en texte, chaînes vides pour les cellules vides)."""
    if is_excel(path):
        yield from _iter_excel_chunks(path, sheet_name, chunk_size)
        return
    fmt = detect_csv_format(path)
    try:
        reader = pd.read_csv(
            path,
            sep=fmt.separator,
            encoding=fmt.encoding,
            dtype=str,
            keep_default_na=False,
            chunksize=chunk_size,
        )
        for chunk in reader:
            chunk.columns = pd.Index(_unique_headers([str(c) for c in chunk.columns]))
            yield chunk
    except (pd.errors.ParserError, UnicodeDecodeError) as exc:
        raise BusinessError(f"Fichier illisible : {exc}") from exc


def _iter_excel_chunks(path: Path, sheet_name: str | None, chunk_size: int) -> Iterator[pd.DataFrame]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook[sheet_name]
            if sheet_name and sheet_name in workbook.sheetnames
            else workbook.worksheets[0]
        )
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            return
        headers = _unique_headers([cell_to_str(v) for v in header_row])
        width = len(headers)
        while block := list(islice(rows, chunk_size)):
            data = [[cell_to_str(v) for v in (list(row) + [None] * width)[:width]] for row in block]
            yield pd.DataFrame(data, columns=headers, dtype=str)
    finally:
        workbook.close()


def preview(path: Path, sheet_name: str | None) -> tuple[list[str], list[list[str]]]:
    """Colonnes et 20 premières lignes."""
    for chunk in iter_chunks(path, sheet_name, PREVIEW_ROWS):
        return list(chunk.columns), chunk.astype(str).values.tolist()
    return [], []
